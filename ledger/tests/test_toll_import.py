import io
from decimal import Decimal
from datetime import date
import openpyxl

from django.test import TestCase, Client
from django.urls import reverse
from django.contrib.auth.models import User, Permission
from django.utils import timezone

from ledger.models import TollAccount, TollRecharge, TollEntry, CompanyAccount, FinancialRecord
from fleet.models import Vehicle
from ledger.services.toll_import import (
    import_fasttag_statement,
    import_bank_statement,
    import_boss_summary,
    normalize_plate,
    build_vehicle_lookup,
)


class TollImportAndReconciliationTests(TestCase):
    def setUp(self):
        # Create users
        self.user = User.objects.create_user(username='toll_admin', password='password123')
        # Assign permissions
        for codename in ['view_tollaccount', 'add_tollentry', 'change_tollentry',
                         'add_tollrecharge', 'change_tollrecharge', 'delete_tollrecharge',
                         'can_view_financial_records']:
            perm = Permission.objects.get(codename=codename)
            self.user.user_permissions.add(perm)

        self.client = Client()
        self.client.login(username='toll_admin', password='password123')

        # Company Accounts
        self.company1 = CompanyAccount.objects.create(
            name='GS KHALSA ROADLINES',
            opening_balance=Decimal('100000.00'),
        )
        self.company2 = CompanyAccount.objects.create(
            name='KHALSA ROADLINES',
            opening_balance=Decimal('50000.00'),
        )

        # Toll Accounts
        self.toll_account_axis = TollAccount.objects.create(
            name='AXIS',
            provider='AXIS BANK',
            opening_balance=Decimal('1000.00'),
            opening_balance_date=date(2026, 4, 1),
            is_default=False
        )
        self.toll_account_livquik = TollAccount.objects.create(
            name='LivQuik',
            provider='LivQuik',
            opening_balance=Decimal('5000.00'),
            opening_balance_date=date(2026, 4, 1),
            is_default=True
        )

        # Vehicles
        self.vehicle1 = Vehicle.objects.create(
            registration_plate='CG 04 PJ 8276',
            make_model='Tata Signa',
            ownership='OWN'
        )
        self.vehicle2 = Vehicle.objects.create(
            registration_plate='MH 40 DC 7782',
            make_model='Ashok Leyland',
            ownership='OWN'
        )
        self.vehicle3 = Vehicle.objects.create(
            registration_plate='CG 04 PX 8214',
            chassis_number='MB1T2VHD8TPML4652',
            make_model='BharatBenz',
            ownership='OWN'
        )

    def test_plate_normalization(self):
        """Test plate normalizer handles spaces, hyphens, and casing."""
        self.assertEqual(normalize_plate('CG 04 PJ 8276'), 'CG04PJ8276')
        self.assertEqual(normalize_plate('cg-04-pj-8276'), 'CG04PJ8276')
        self.assertEqual(normalize_plate('MH 40 DC 7782'), 'MH40DC7782')

    def test_build_vehicle_lookup_includes_chassis(self):
        """Vehicle lookup maps both registration plate and chassis number."""
        lookup = build_vehicle_lookup()
        self.assertEqual(lookup.get('CG04PJ8276'), self.vehicle1)
        self.assertEqual(lookup.get('MH40DC7782'), self.vehicle2)
        self.assertEqual(lookup.get('CG04PX8214'), self.vehicle3)
        self.assertEqual(lookup.get('MB1T2VHD8TPML4652'), self.vehicle3)

    def test_import_bank_statement_matching_by_chassis(self):
        """Statement containing chassis number in Truck Number column matches Vehicle by chassis."""
        wb = openpyxl.Workbook()
        ws = wb.active
        for _ in range(7):
            ws.append([None] * 9)
        ws.append(['Transaction Time', 'Nature (C/D)', 'Amount', 'Description',
                   'Truck Number', 'Transaction ID', 'Opening Balance', 'Closing Balance', 'Vendor'])
        ws.append(['21 Oct 26 07:05 AM', 'Debit', 490.0, 'FasTag Toll Payment at Khed - Maharashtra',
                   'MB1T2VHD8TPML4652', 'TXN-CHASSIS-001', 5000.0, 4510.0, 'LIVQUIK'])

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        summary = import_bank_statement(buf, toll_account=self.toll_account_livquik, user=self.user)
        self.assertEqual(summary.toll_entries_created, 1)
        self.assertEqual(len(summary.unmatched_vehicles), 0)

        entry = TollEntry.objects.get(transaction_reference='TXN-CHASSIS-001')
        self.assertEqual(entry.vehicle, self.vehicle3)
        self.assertIn('Tag Chassis: MB1T2VHD8TPML4652', entry.description)

    def test_unallocated_recharge_creates_financial_record_with_null_account(self):
        """An unallocated TollRecharge creates FinancialRecord with account=None so reports account for it."""
        initial_toll_balance = self.toll_account_livquik.current_balance_cached
        recharge = TollRecharge.objects.create(
            toll_account=self.toll_account_livquik,
            company_account=None,
            date=date(2026, 4, 2),
            amount=Decimal('70000.00'),
            payment_mode=TollRecharge.MODE_NET_BANKING,
            reference_number='TOPUP-001',
            created_by=self.user
        )

        # Toll balance refreshed
        self.toll_account_livquik.refresh_from_db()
        self.assertEqual(
            self.toll_account_livquik.current_balance_cached,
            initial_toll_balance + Decimal('70000.00')
        )

        # FinancialRecord exists with account=None
        fin_rec = FinancialRecord.objects.filter(associated_toll_recharge=recharge).first()
        self.assertIsNotNone(fin_rec)
        self.assertIsNone(fin_rec.account)
        self.assertEqual(fin_rec.amount, Decimal('70000.00'))
        self.assertEqual(fin_rec.category.name, 'Toll Recharge')

    def test_reconcile_and_unlink_recharge(self):
        """Assigning a firm sets FinancialRecord.account; unlinking reverts account to None."""
        recharge = TollRecharge.objects.create(
            toll_account=self.toll_account_livquik,
            company_account=None,
            date=date(2026, 4, 2),
            amount=Decimal('50000.00'),
            payment_mode=TollRecharge.MODE_NET_BANKING,
            reference_number='TOPUP-002',
            created_by=self.user
        )

        # Post to reconcile
        response = self.client.post(reverse('toll-reconcile'), {
            'action': 'reconcile_single',
            'recharge_id': recharge.pk,
            'company_account': self.company1.pk,
        })
        self.assertRedirects(response, reverse('toll-reconcile'))

        recharge.refresh_from_db()
        self.assertEqual(recharge.company_account, self.company1)

        # FinancialRecord updated to company1
        fin_record = FinancialRecord.objects.filter(associated_toll_recharge=recharge).first()
        self.assertIsNotNone(fin_record)
        self.assertEqual(fin_record.account, self.company1)
        self.assertEqual(fin_record.amount, Decimal('50000.00'))

        # Now unlink
        unlink_response = self.client.get(reverse('toll-recharge-unlink', kwargs={'pk': recharge.pk}))
        self.assertRedirects(unlink_response, reverse('toll-reconcile'))

        recharge.refresh_from_db()
        self.assertIsNone(recharge.company_account)
        # FinancialRecord preserved with account=None
        fin_record.refresh_from_db()
        self.assertIsNone(fin_record.account)

    def test_financial_summary_includes_unallocated_toll_recharges(self):
        """Financial summary report view accounts for unallocated toll recharges as real cash outflow."""
        TollRecharge.objects.create(
            toll_account=self.toll_account_livquik,
            company_account=None,
            date=date(2026, 4, 15),
            amount=Decimal('45000.00'),
            payment_mode=TollRecharge.MODE_NET_BANKING,
            reference_number='TOPUP-APR-001',
            created_by=self.user
        )

        response = self.client.get(reverse('financial-summary') + '?month=4&year=2026')
        self.assertEqual(response.status_code, 200)
        self.assertGreaterEqual(response.context['monthly_expenses'], Decimal('45000.00'))
        self.assertGreaterEqual(response.context['yearly_expenses'], Decimal('45000.00'))
        self.assertGreaterEqual(response.context['unallocated_monthly_expenses'], Decimal('45000.00'))

        cat_names = [c['name'] for c in response.context['category_breakdown']]
        self.assertIn('Toll Recharge', cat_names)

    def test_bulk_reconciliation(self):
        """Bulk reconciling assigns multiple unallocated recharges to a company account."""
        r1 = TollRecharge.objects.create(
            toll_account=self.toll_account_livquik,
            company_account=None,
            date=date(2026, 4, 2),
            amount=Decimal('20000.00'),
            reference_number='BULK-001'
        )
        r2 = TollRecharge.objects.create(
            toll_account=self.toll_account_livquik,
            company_account=None,
            date=date(2026, 4, 3),
            amount=Decimal('30000.00'),
            reference_number='BULK-002'
        )

        response = self.client.post(reverse('toll-reconcile'), {
            'action': 'reconcile_bulk',
            'recharge_ids': [r1.pk, r2.pk],
            'bulk_company_account': self.company2.pk,
        })
        self.assertRedirects(response, reverse('toll-reconcile'))

        r1.refresh_from_db()
        r2.refresh_from_db()
        self.assertEqual(r1.company_account, self.company2)
        self.assertEqual(r2.company_account, self.company2)
        self.assertEqual(FinancialRecord.objects.filter(associated_toll_recharge__in=[r1, r2]).count(), 2)

    def test_import_bank_statement_debits(self):
        """Verify importing bank spreadsheet creates TollEntry records and skips duplicates."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Account_Summary"

        # Rows 1-7 metadata
        for _ in range(7):
            ws.append([None] * 9)

        # Row 8: Header
        ws.append(['Transaction Time', 'Nature (C/D)', 'Amount', 'Description',
                   'Truck Number', 'Transaction ID', 'Opening Balance', 'Closing Balance', 'Vendor'])

        # Row 9: Credit (internal sweep, should be skipped)
        ws.append(['21 Oct 26 07:09 AM', 'Credit', 375.0, 'FasTag Recharge', '', '174899355', 175.0, 550.0, 'AXIS'])

        # Row 10: Debit (valid toll payment for vehicle1)
        ws.append(['21 Oct 26 07:05 AM', 'Debit', 375.0, 'FasTag Toll Payment at Mathni - Maharashtra',
                   'CG04PJ8276', 'TXN-BANK-001', 550.0, 175.0, 'AXIS'])

        # Row 11: Debit (valid toll payment for vehicle2)
        ws.append(['20 Oct 26 05:38 PM', 'Debit', 440.0, 'FasTag Toll Payment at Arambha',
                   'MH40DC7782', 'TXN-BANK-002', 550.0, 110.0, 'AXIS'])

        # Save to memory buffer
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        summary = import_bank_statement(buf, toll_account=self.toll_account_axis, user=self.user)

        self.assertEqual(summary.toll_entries_created, 2)
        self.assertEqual(summary.total_debit_amount, Decimal('815.00'))
        self.assertEqual(TollEntry.objects.count(), 2)

        entry1 = TollEntry.objects.get(transaction_reference='TXN-BANK-001')
        self.assertEqual(entry1.vehicle, self.vehicle1)
        self.assertEqual(entry1.amount, Decimal('375.00'))
        self.assertEqual(entry1.toll_plaza, 'Mathni - Maharashtra')

        # Re-importing same buffer should skip duplicates
        buf.seek(0)
        summary2 = import_bank_statement(buf, toll_account=self.toll_account_axis, user=self.user)
        self.assertEqual(summary2.toll_entries_created, 0)
        self.assertEqual(summary2.duplicate_entries_skipped, 2)

    def test_import_boss_summary_credits(self):
        """Verify importing Boss summary creates unallocated TollRecharge records."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Boss_Account_Summary"

        # Rows 1-20 metadata
        for _ in range(20):
            ws.append([None] * 9)

        # Row 21: Header
        ws.append(['Transaction Time', 'Nature (C/D)', 'Amount', 'From/To',
                   'Description', 'Truck Number', 'TransactionId', 'Opening Balance', 'Closing Balance'])

        # Row 22: Debit in Boss (internal tag recharge, should be skipped)
        ws.append(['1 Apr 26 01:31 AM', 'Debit', 440, 'Wallet', 'FasTag Recharge', '', '1174375260', 57232.61, 56792.61])

        # Row 23: Credit (Net Banking TopUp, should be imported as unallocated recharge)
        ws.append(['2 Apr 26 04:49 AM', 'Credit', 70000, 'Net Banking', 'TopUp', '', 'BOSS-TXN-001', 28004.61, 98004.61])

        # Row 24: Credit (UPI TopUp)
        ws.append(['3 Apr 26 04:54 PM', 'Credit', 50000, 'UPI', 'TopUp', '', 'BOSS-TXN-002', 37288.61, 87288.61])

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        summary = import_boss_summary(buf, toll_account=self.toll_account_livquik, user=self.user)

        self.assertEqual(summary.recharges_created, 2)
        self.assertEqual(summary.total_credit_amount, Decimal('120000.00'))

        recharge = TollRecharge.objects.get(reference_number='BOSS-TXN-001')
        self.assertIsNone(recharge.company_account)
        self.assertEqual(recharge.amount, Decimal('70000.00'))
        self.assertEqual(recharge.payment_mode, TollRecharge.MODE_NET_BANKING)

        # Re-importing should skip duplicates
        buf.seek(0)
        summary2 = import_boss_summary(buf, toll_account=self.toll_account_livquik, user=self.user)
        self.assertEqual(summary2.recharges_created, 0)
        self.assertEqual(summary2.duplicate_recharges_skipped, 2)

    def test_import_view_get_and_post(self):
        """Test TollStatementImportView renders GET and handles POST upload."""
        # 1. GET
        response = self.client.get(reverse('toll-statement-import'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Import FASTag Statement')

        # 2. POST with simulated file
        wb = openpyxl.Workbook()
        ws = wb.active
        for _ in range(7):
            ws.append([None] * 9)
        ws.append(['Transaction Time', 'Nature (C/D)', 'Amount', 'Description',
                   'Truck Number', 'Transaction ID', 'Opening Balance', 'Closing Balance', 'Vendor'])
        ws.append(['21 Oct 26 07:05 AM', 'Debit', 375.0, 'FasTag Toll Payment at Mathni',
                   'CG04PJ8276', 'VIEW-TXN-001', 550.0, 175.0, 'AXIS'])

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        buf.name = 'test_axis.xlsx'

        from django.core.files.uploadedfile import SimpleUploadedFile
        uploaded_file = SimpleUploadedFile('test_axis.xlsx', buf.read(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

        post_resp = self.client.post(reverse('toll-statement-import'), {
            'file': uploaded_file,
            'toll_account': self.toll_account_axis.pk,
        })
        self.assertRedirects(post_resp, reverse('toll-ledger'))
        self.assertTrue(TollEntry.objects.filter(transaction_reference='VIEW-TXN-001').exists())

    def test_toll_ledger_renders_with_unallocated_badge(self):
        """Test TollLedgerView displays unallocated recharges and badge."""
        TollRecharge.objects.create(
            toll_account=self.toll_account_livquik,
            company_account=None,
            date=date(2026, 4, 2),
            amount=Decimal('70000.00'),
            reference_number='PENDING-001'
        )

        response = self.client.get(reverse('toll-ledger'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Unallocated')
        self.assertContains(response, 'Reconcile')

